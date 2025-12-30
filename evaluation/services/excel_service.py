"""
Excel Service
=============
Handles reading and writing to Excel evaluation files.
Follows Single Responsibility: Only handles Excel I/O.
"""

import os
from typing import List, Optional, Dict, Any
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from ..models import Question, BotResponse, EvaluationResult
from ..config import EvaluationConfig, BotConfig, DEFAULT_QUESTIONS


class ExcelService:
    """
    Service for managing Excel evaluation files.
    
    Responsibilities:
    - Create evaluation sheets
    - Read questions from sheets
    - Write responses and evaluations
    - Generate reports
    """
    
    def __init__(self, config: EvaluationConfig):
        self.config = config
        self.workbook = None
        self.worksheet = None
        self._file_path = config.excel_file
    
    @property
    def file_path(self) -> str:
        """Get the Excel file path."""
        return self._file_path
    
    def load_or_create(self, file_path: Optional[str] = None) -> bool:
        """
        Load existing workbook or create a new one.
        
        Args:
            file_path: Optional path to Excel file
            
        Returns:
            True if successful
        """
        if file_path:
            self._file_path = file_path
        
        try:
            self.workbook = openpyxl.load_workbook(self._file_path)
            self.worksheet = self.workbook.active
            return True
        except FileNotFoundError:
            self.create_new_sheet()
            return True
        except Exception as e:
            print(f"[ERROR] Failed to load Excel: {e}")
            return False
    
    def create_new_sheet(
        self,
        questions: Optional[List[str]] = None,
        file_path: Optional[str] = None
    ) -> str:
        """
        Create a new evaluation Excel sheet.
        
        Args:
            questions: Optional list of questions (uses defaults if None)
            file_path: Optional file path
            
        Returns:
            Path to the created file
        """
        if file_path:
            self._file_path = file_path
        
        questions = questions or DEFAULT_QUESTIONS
        
        # Create workbook
        self.workbook = openpyxl.Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "Bot Evaluation"
        
        # Styles
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Headers
        headers = ["Question"]
        for bot_name, bot_config in self.config.bots.items():
            headers.append(bot_name)
            headers.append(f"{bot_name} Eval")
            headers.append(f"{bot_name} Comment")
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = self.worksheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Column widths
        self.worksheet.column_dimensions['A'].width = 50
        col_idx = 2
        for bot_config in self.config.bots.values():
            # Response column
            self.worksheet.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx)
            ].width = 35
            # Eval column
            self.worksheet.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx + 1)
            ].width = 20
            # Comment column
            self.worksheet.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx + 2)
            ].width = 30
            col_idx += 3
        
        # Write questions and format cells
        for row_idx, question in enumerate(questions, 2):
            # Question cell
            cell = self.worksheet.cell(row=row_idx, column=1, value=question)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border
            
            # Bot columns
            col_idx = 2
            for bot_config in self.config.bots.values():
                # Response column
                response_cell = self.worksheet.cell(row=row_idx, column=col_idx, value="")
                response_cell.fill = PatternFill(
                    start_color=bot_config.color_response,
                    end_color=bot_config.color_response,
                    fill_type="solid"
                )
                response_cell.border = thin_border
                response_cell.alignment = Alignment(vertical="center", wrap_text=True)
                
                # Eval column
                eval_cell = self.worksheet.cell(row=row_idx, column=col_idx + 1, value="")
                eval_cell.fill = PatternFill(
                    start_color=bot_config.color_eval,
                    end_color=bot_config.color_eval,
                    fill_type="solid"
                )
                eval_cell.border = thin_border
                eval_cell.alignment = Alignment(vertical="center", wrap_text=True)
                
                # Comment column
                comment_cell = self.worksheet.cell(row=row_idx, column=col_idx + 2, value="")
                comment_cell.fill = PatternFill(
                    start_color=bot_config.color_comment,
                    end_color=bot_config.color_comment,
                    fill_type="solid"
                )
                comment_cell.border = thin_border
                comment_cell.alignment = Alignment(vertical="center", wrap_text=True)
                
                col_idx += 3
        
        # Add empty rows
        for row_idx in range(len(questions) + 2, len(questions) + 12):
            for col in range(1, len(headers) + 1):
                cell = self.worksheet.cell(row=row_idx, column=col, value="")
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        
        # Freeze header
        self.worksheet.freeze_panes = 'A2'
        self.worksheet.row_dimensions[1].height = 30
        
        # Save
        self.save()
        print(f"[OK] Created evaluation sheet: {self._file_path}")
        
        return self._file_path
    
    def get_questions(self) -> List[Question]:
        """
        Read all questions from the Excel sheet.
        
        Returns:
            List of Question objects
        """
        if not self.worksheet:
            self.load_or_create()
        
        questions = []
        row = 2  # Start after header
        
        while True:
            cell = self.worksheet[f"A{row}"]
            if cell.value is None or str(cell.value).strip() == "":
                break
            questions.append(Question(
                index=row - 2,  # 0-based index
                text=str(cell.value)
            ))
            row += 1
        
        return questions
    
    def write_response(self, response: BotResponse) -> bool:
        """
        Write a bot's response to the Excel sheet.
        
        Args:
            response: The response to write
            
        Returns:
            True if successful
        """
        if not self.worksheet:
            self.load_or_create()
        
        bot_config = self.config.get_bot(response.bot_name)
        if not bot_config:
            print(f"[ERROR] Unknown bot: {response.bot_name}")
            return False
        
        row = response.question_index + 2  # +2 for header and 0-based index
        col = bot_config.response_column
        
        cell = self.worksheet[f"{col}{row}"]
        cell.value = response.response_text
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        
        return True
    
    def write_evaluation(self, result: EvaluationResult) -> bool:
        """
        Write an evaluation result to the Excel sheet.
        
        Args:
            result: The evaluation result to write
            
        Returns:
            True if successful
        """
        if not self.worksheet:
            self.load_or_create()
        
        bot_config = self.config.get_bot(result.bot_name)
        if not bot_config:
            print(f"[ERROR] Unknown bot: {result.bot_name}")
            return False
        
        row = result.question_index + 2
        col = bot_config.eval_column
        
        # Format evaluation text
        if result.score is not None:
            eval_text = f"{result.score}/10 - {result.evaluation_text}"
        else:
            eval_text = result.evaluation_text
        
        cell = self.worksheet[f"{col}{row}"]
        cell.value = eval_text
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        
        return True
    
    def write_comment(self, bot_name: str, question_index: int, comment: str) -> bool:
        """
        Write a comment for a bot's response.
        
        Args:
            bot_name: Name of the bot
            question_index: Index of the question (0-based)
            comment: The comment text to write
            
        Returns:
            True if successful
        """
        if not self.worksheet:
            self.load_or_create()
        
        bot_config = self.config.get_bot(bot_name)
        if not bot_config:
            print(f"[ERROR] Unknown bot: {bot_name}")
            return False
        
        row = question_index + 2
        col = bot_config.comment_column
        
        cell = self.worksheet[f"{col}{row}"]
        cell.value = comment
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        
        return True
    
    def save(self) -> bool:
        """Save the workbook."""
        try:
            self.workbook.save(self._file_path)
            return True
        except PermissionError:
            print(f"[ERROR] Cannot save - file is open: {self._file_path}")
            return False
        except Exception as e:
            print(f"[ERROR] Failed to save: {e}")
            return False
    
    def get_summary(self) -> Dict[str, Any]:
        """
        Get a summary of the current evaluation state.
        
        Returns:
            Dictionary with summary statistics
        """
        if not self.worksheet:
            self.load_or_create()
        
        questions = self.get_questions()
        summary = {
            "file": self._file_path,
            "total_questions": len(questions),
            "bots": {}
        }
        
        for bot_name, bot_config in self.config.bots.items():
            filled_responses = 0
            filled_evals = 0
            
            for q in questions:
                row = q.index + 2
                
                response_cell = self.worksheet[f"{bot_config.response_column}{row}"]
                if response_cell.value and str(response_cell.value).strip():
                    filled_responses += 1
                
                eval_cell = self.worksheet[f"{bot_config.eval_column}{row}"]
                if eval_cell.value and str(eval_cell.value).strip():
                    filled_evals += 1
            
            summary["bots"][bot_name] = {
                "assistant_id": bot_config.assistant_id[:20] + "..." if bot_config.assistant_id else "Not configured",
                "responses_filled": filled_responses,
                "evaluations_filled": filled_evals,
                "total": len(questions)
            }
        
        return summary


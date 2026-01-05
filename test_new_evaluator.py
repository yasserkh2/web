"""
Test the new segment-based evaluator with 3 segments.
Uses GPT-4o-mini and saves results to JSON and Excel.
"""

import os
import json
import time
from datetime import datetime
from pathlib import Path
import sys

# Fix Windows console encoding
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

# Add project root to path
sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from evaluation.config import LLMConfig, BotConfig
from evaluation.evaluators.llm_evaluator import LLMEvaluator
from evaluation.models import BotResponse, Question

# Delay between API calls to avoid rate limits
API_DELAY_SECONDS = 25


def load_dataset(segment_name: str) -> dict:
    """Load evaluation dataset for a segment."""
    filename = segment_name.lower().replace(" ", "_") + ".json"
    dataset_path = Path("datasets") / filename
    
    if dataset_path.exists():
        with open(dataset_path, "r", encoding="utf-8") as f:
            return json.load(f)
    return None


def run_evaluation(segment_name: str, evaluator: LLMEvaluator, max_cases: int = 3) -> dict:
    """Run evaluation for a segment and return results."""
    print(f"\n{'='*60}")
    print(f"Evaluating: {segment_name}")
    print(f"{'='*60}")
    
    # Load dataset
    dataset = load_dataset(segment_name)
    if not dataset:
        print(f"  [ERROR] Dataset not found for {segment_name}")
        return None
    
    # Create bot config
    bot_config = BotConfig(
        name=segment_name,
        assistant_id="",
        description=f"Physician persona: {segment_name}"
    )
    
    results = {
        "segment": segment_name,
        "timestamp": datetime.now().isoformat(),
        "model": evaluator.config.model,
        "total_evaluated": 0,
        "average_score": 0.0,
        "low_score_count": 0,
        "pass_rate": 0.0,
        "prompt_improvement_suggestions": [],
        "detailed_results": []
    }
    
    cases = dataset.get("cases", [])[:max_cases]
    total_score = 0
    
    for i, case in enumerate(cases):
        question_text = case.get("question", case.get("input", ""))
        response_text = case.get("expected_response", case.get("ground_truth", ""))
        
        print(f"\n  Case {i+1}/{len(cases)}: {question_text[:50]}...")
        
        # Create evaluation objects
        question = Question(index=case.get("id", i+1), text=question_text)
        bot_response = BotResponse(
            bot_name=segment_name,
            question_index=case.get("id", i+1),
            question_text=question_text,
            response_text=response_text
        )
        
        # Run evaluation
        try:
            eval_result = evaluator.evaluate(bot_response, bot_config, question)
            
            # Extract data from criteria_scores
            raw_eval = eval_result.criteria_scores.get("_raw_evaluation", {})
            segment_fit_score = raw_eval.get("segment_fit_score", eval_result.score / 2 if eval_result.score else 0)
            
            # Build result entry
            result_entry = {
                "score": segment_fit_score,
                "segment_fit_score": segment_fit_score,
                "supporting_quotes": raw_eval.get("supporting_quotes", []),
                "segment_leaks": raw_eval.get("segment_leaks", []),
                "rewrite_if_needed": raw_eval.get("rewrite_if_needed", ""),
                "comment": eval_result.evaluation_text[:500] if eval_result.evaluation_text else "",
                "question_id": case.get("id", i+1),
                "question": question_text,
                "response_preview": response_text[:200] + "..." if len(response_text) > 200 else response_text
            }
            
            results["detailed_results"].append(result_entry)
            total_score += segment_fit_score
            
            if segment_fit_score <= 3:
                results["low_score_count"] += 1
            
            print(f"    Score: {segment_fit_score}/5")
            print(f"    Leaks: {len(raw_eval.get('segment_leaks', []))}")
            
            # Delay to avoid rate limits
            if i < len(cases) - 1:
                print(f"    Waiting {API_DELAY_SECONDS}s for rate limit...")
                time.sleep(API_DELAY_SECONDS)
            
        except Exception as e:
            print(f"    [ERROR] {str(e)}")
            results["detailed_results"].append({
                "score": 0,
                "error": str(e),
                "question_id": case.get("id", i+1),
                "question": question_text
            })
    
    # Calculate summary stats
    results["total_evaluated"] = len(results["detailed_results"])
    if results["total_evaluated"] > 0:
        results["average_score"] = round(total_score / results["total_evaluated"], 2)
        results["pass_rate"] = round((results["total_evaluated"] - results["low_score_count"]) / results["total_evaluated"] * 100, 1)
    
    print(f"\n  Summary: Avg Score = {results['average_score']}/5, Pass Rate = {results['pass_rate']}%")
    
    return results


def save_json_results(results: dict, segment_name: str):
    """Save results to JSON file."""
    output_dir = Path("eval_results")
    output_dir.mkdir(exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{segment_name.lower().replace(' ', '_')}_eval_{timestamp}.json"
    output_path = output_dir / filename
    
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2, ensure_ascii=False)
    
    print(f"  Saved: {output_path}")
    return output_path


def save_to_excel(all_results: list, excel_path: str = "bot_evaluation_with_comments.xlsx"):
    """Save all results to Excel file."""
    # Create new workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation Results"
    
    # Header style
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    # Headers
    headers = ["Segment", "Question ID", "Question", "Score (0-5)", "Supporting Quotes", "Segment Leaks", "Rewrite Needed", "Response Preview"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Data rows
    row = 2
    for segment_results in all_results:
        segment = segment_results["segment"]
        for detail in segment_results["detailed_results"]:
            ws.cell(row=row, column=1, value=segment)
            ws.cell(row=row, column=2, value=detail.get("question_id", ""))
            ws.cell(row=row, column=3, value=detail.get("question", ""))
            ws.cell(row=row, column=4, value=detail.get("score", 0))
            
            # Format supporting quotes
            quotes = detail.get("supporting_quotes", [])
            quotes_text = "\n".join([f'"{q.get("quote", "")}" → {q.get("trait_matched", "")}' for q in quotes])
            ws.cell(row=row, column=5, value=quotes_text)
            
            # Format segment leaks
            leaks = detail.get("segment_leaks", [])
            leaks_text = "\n".join([f'{l.get("leak_text", "")} → {l.get("resembles_segment", "")}' for l in leaks])
            ws.cell(row=row, column=6, value=leaks_text)
            
            ws.cell(row=row, column=7, value=detail.get("rewrite_if_needed", ""))
            ws.cell(row=row, column=8, value=detail.get("response_preview", ""))
            
            # Apply wrap text to all cells
            for col in range(1, 9):
                ws.cell(row=row, column=col).alignment = Alignment(vertical="top", wrap_text=True)
            
            row += 1
    
    # Add summary sheet
    ws_summary = wb.create_sheet("Summary")
    ws_summary.cell(row=1, column=1, value="Segment").font = Font(bold=True)
    ws_summary.cell(row=1, column=2, value="Total Evaluated").font = Font(bold=True)
    ws_summary.cell(row=1, column=3, value="Average Score").font = Font(bold=True)
    ws_summary.cell(row=1, column=4, value="Pass Rate").font = Font(bold=True)
    ws_summary.cell(row=1, column=5, value="Low Scores").font = Font(bold=True)
    
    for i, result in enumerate(all_results, 2):
        ws_summary.cell(row=i, column=1, value=result["segment"])
        ws_summary.cell(row=i, column=2, value=result["total_evaluated"])
        ws_summary.cell(row=i, column=3, value=result["average_score"])
        ws_summary.cell(row=i, column=4, value=f"{result['pass_rate']}%")
        ws_summary.cell(row=i, column=5, value=result["low_score_count"])
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 50
    ws.column_dimensions['G'].width = 50
    ws.column_dimensions['H'].width = 50
    
    wb.save(excel_path)
    print(f"\n[OK] Saved Excel: {excel_path}")


def main():
    # Get API key from environment
    api_key = os.environ.get("OPENAI_API_KEY", "")
    if not api_key:
        # Try to load from .env file
        env_path = Path(".env")
        if env_path.exists():
            with open(env_path) as f:
                for line in f:
                    if line.startswith("OPENAI_API_KEY="):
                        api_key = line.split("=", 1)[1].strip().strip('"').strip("'")
                        break
    
    if not api_key:
        print("ERROR: OPENAI_API_KEY not found in environment or .env file")
        print("Please set: $env:OPENAI_API_KEY = 'your-key-here'")
        return 1
    
    # Create evaluator with GPT-4o-mini
    config = LLMConfig(
        provider="openai",
        model="gpt-4o-mini",
        api_key=api_key,
        temperature=0.3
    )
    evaluator = LLMEvaluator(config)
    
    print("=" * 60)
    print("NEW SEGMENT EVALUATOR TEST")
    print(f"Model: {config.model}")
    print("=" * 60)
    
    # Test 3 segments
    segments = ["The Innovator", "The Traditionalist", "The Evidence Purist"]
    all_results = []
    
    for segment in segments:
        results = run_evaluation(segment, evaluator, max_cases=3)
        if results:
            save_json_results(results, segment)
            all_results.append(results)
    
    # Save to Excel
    if all_results:
        save_to_excel(all_results)
    
    print("\n" + "=" * 60)
    print("EVALUATION COMPLETE")
    print("=" * 60)
    
    return 0


if __name__ == "__main__":
    sys.exit(main())


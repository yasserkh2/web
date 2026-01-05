#!/usr/bin/env python
"""
Evaluation Agent - Direct Excel Mode
=====================================
Reads questions and responses from Excel, evaluates, writes results back.

Usage:
    python eval_agent.py --model gpt-4o-mini
"""

import os
import sys
import time
import argparse
from datetime import datetime
from pathlib import Path

# Fix Windows console encoding
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except:
        pass

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

from evaluation.config import LLMConfig, BotConfig
from evaluation.evaluators.llm_evaluator import LLMEvaluator
from evaluation.models import BotResponse, Question

# Excel file path
EXCEL_FILE = "bot_evaluation_with_comments.xlsx"

# Segment columns mapping (adjust based on your Excel structure)
# Format: segment_name -> (response_col, eval_col, comment_col)
# 7 columns per segment: Response, Eval, Comment, Eval_Aboubakr, Comment_Aboubakr, Eval_Thomas, Comment_Thomas
SEGMENT_COLUMNS = {
    "The Traditionalist": ("B", "C", "D"),   # Response=B, Eval=C, Comment=D
    "The Innovator": ("I", "J", "K"),         # Response=I, Eval=J, Comment=K
    "The Evidence Purist": ("P", "Q", "R"),   # Response=P, Eval=Q, Comment=R
}

# Expected headers for each column
# A=Question, then for each segment: Response, Eval, Comment, Eval_Aboubakr, Comment_Aboubakr, Eval_Thomas, Comment_Thomas
COLUMN_HEADERS = {
    "A": "Question",
    # The Traditionalist (B-H)
    "B": "The Traditionalist",
    "C": "The Traditionalist Eval",
    "D": "The Traditionalist Comment",
    "E": "The Traditionalist Eval_Aboubakr",
    "F": "The Traditionalist Comment_Aboubakr",
    "G": "The Traditionalist Eval_Thomas",
    "H": "The Traditionalist Comment_Thomas",
    # The Innovator (I-O)
    "I": "The Innovator",
    "J": "The Innovator Eval",
    "K": "The Innovator Comment",
    "L": "The Innovator Eval_Aboubakr",
    "M": "The Innovator Comment_Aboubakr",
    "N": "The Innovator Eval_Thomas",
    "O": "The Innovator Comment_Thomas",
    # The Evidence Purist (P-V)
    "P": "The Evidence Purist",
    "Q": "The Evidence Purist Eval",
    "R": "The Evidence Purist Comment",
    "S": "The Evidence Purist Eval_Aboubakr",
    "T": "The Evidence Purist Comment_Aboubakr",
    "U": "The Evidence Purist Eval_Thomas",
    "V": "The Evidence Purist Comment_Thomas",
}

# API delay between calls (seconds)
API_DELAY = 30


def get_column_letter(col_idx):
    """Convert column index to letter (1=A, 2=B, etc.)"""
    result = ""
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = chr(65 + remainder) + result
    return result


def run_evaluation(excel_path: str, evaluator: LLMEvaluator, segments: list = None, start_row: int = 2, args=None):
    """
    Read Excel, evaluate responses, write results back.
    
    Args:
        excel_path: Path to Excel file
        evaluator: LLM evaluator instance
        segments: List of segments to evaluate (None = all)
        start_row: First data row (default 2, assuming row 1 is header)
        args: Command line arguments
    """
    # Create a simple namespace if args not provided
    if args is None:
        class Args:
            force = False
        args = Args()
    print(f"\nLoading: {excel_path}")
    
    try:
        wb = load_workbook(excel_path)
        ws = wb.active
    except Exception as e:
        print(f"ERROR: Cannot open Excel file: {e}")
        return
    
    # Ensure headers are properly set in row 1
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col_letter, header_text in COLUMN_HEADERS.items():
        cell = ws[f"{col_letter}1"]
        if cell.value != header_text:
            cell.value = header_text
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
    
    # Set column widths
    ws.column_dimensions['A'].width = 50  # Question
    # The Traditionalist (B-H)
    ws.column_dimensions['B'].width = 50  # Response
    ws.column_dimensions['C'].width = 12  # Eval
    ws.column_dimensions['D'].width = 40  # Comment
    ws.column_dimensions['E'].width = 12  # Eval_Aboubakr
    ws.column_dimensions['F'].width = 40  # Comment_Aboubakr
    ws.column_dimensions['G'].width = 12  # Eval_Thomas
    ws.column_dimensions['H'].width = 40  # Comment_Thomas
    # The Innovator (I-O)
    ws.column_dimensions['I'].width = 50  # Response
    ws.column_dimensions['J'].width = 12  # Eval
    ws.column_dimensions['K'].width = 40  # Comment
    ws.column_dimensions['L'].width = 12  # Eval_Aboubakr
    ws.column_dimensions['M'].width = 40  # Comment_Aboubakr
    ws.column_dimensions['N'].width = 12  # Eval_Thomas
    ws.column_dimensions['O'].width = 40  # Comment_Thomas
    # The Evidence Purist (P-V)
    ws.column_dimensions['P'].width = 50  # Response
    ws.column_dimensions['Q'].width = 12  # Eval
    ws.column_dimensions['R'].width = 40  # Comment
    ws.column_dimensions['S'].width = 12  # Eval_Aboubakr
    ws.column_dimensions['T'].width = 40  # Comment_Aboubakr
    ws.column_dimensions['U'].width = 12  # Eval_Thomas
    ws.column_dimensions['V'].width = 40  # Comment_Thomas
    
    wb.save(excel_path)
    print("Headers verified/updated")
    
    # Get max row
    max_row = ws.max_row
    print(f"Found {max_row - start_row + 1} rows to process")
    
    # Determine which segments to evaluate
    segments_to_eval = segments or list(SEGMENT_COLUMNS.keys())
    
    total_evaluations = 0
    
    for row in range(start_row, max_row + 1):
        # Get question from column A
        question_text = ws[f"A{row}"].value
        if not question_text:
            continue
        
        print(f"\n[Row {row}] {str(question_text)[:50]}...")
        
        for segment in segments_to_eval:
            if segment not in SEGMENT_COLUMNS:
                print(f"  WARNING: Unknown segment '{segment}'")
                continue
            
            response_col, eval_col, comment_col = SEGMENT_COLUMNS[segment]
            
            # Get response from Excel
            response_text = ws[f"{response_col}{row}"].value
            if not response_text:
                print(f"  {segment}: No response, skipping")
                continue
            
            # Check if already evaluated (skip unless --force)
            existing_eval = ws[f"{eval_col}{row}"].value
            if existing_eval is not None and existing_eval != "" and not args.force:
                print(f"  {segment}: Already evaluated ({existing_eval}/5), skipping")
                continue
            
            print(f"  {segment}: Evaluating...")
            
            # Create evaluation objects
            bot_config = BotConfig(name=segment, assistant_id="", description=f"Physician: {segment}")
            question = Question(index=row, text=str(question_text))
            bot_response = BotResponse(
                bot_name=segment,
                question_index=row,
                question_text=str(question_text),
                response_text=str(response_text)
            )
            
            try:
                # Run evaluation
                eval_result = evaluator.evaluate(bot_response, bot_config, question)
                
                score = eval_result.score if eval_result.score is not None else 0
                comment = eval_result.evaluation_text or ""
                
                # Write score to Excel
                ws[f"{eval_col}{row}"] = score
                
                # Write comment only if score <= 3
                if score <= 3:
                    ws[f"{comment_col}{row}"] = comment
                else:
                    ws[f"{comment_col}{row}"] = ""
                
                print(f"    Score: {score}/5")
                if score <= 3 and comment:
                    print(f"    Comment: {comment[:60]}...")
                
                total_evaluations += 1
                
                # Save after each evaluation (in case of crash)
                wb.save(excel_path)
                
                # Delay for rate limits
                print(f"    [Waiting {API_DELAY}s...]")
                time.sleep(API_DELAY)
                
            except Exception as e:
                error_msg = str(e)
                print(f"    ERROR: {error_msg[:80]}")
                
                # Check for rate limit
                if "429" in error_msg or "rate_limit" in error_msg.lower():
                    print(f"\n    *** RATE LIMIT - waiting 60s ***")
                    time.sleep(60)
                else:
                    ws[f"{eval_col}{row}"] = "ERROR"
                    ws[f"{comment_col}{row}"] = error_msg[:200]
                    wb.save(excel_path)
    
    # Final save
    wb.save(excel_path)
    print(f"\n[OK] Saved {total_evaluations} evaluations to {excel_path}")


def main():
    parser = argparse.ArgumentParser(description="Evaluate responses in Excel file")
    parser.add_argument("--model", type=str, default="gpt-4o-mini", help="Model to use")
    parser.add_argument("--file", type=str, default=EXCEL_FILE, help="Excel file path")
    parser.add_argument("--segment", type=str, help="Evaluate only this segment")
    parser.add_argument("--delay", type=int, default=30, help="Delay between API calls")
    parser.add_argument("--start-row", type=int, default=2, help="First data row")
    parser.add_argument("--force", action="store_true", help="Re-evaluate already evaluated rows")
    
    args = parser.parse_args()
    
    global API_DELAY
    API_DELAY = args.delay
    
    # Get API key
    api_key = os.environ.get("OPENAI_API_KEY", "")
    if not api_key:
        env_path = Path(".env")
        if env_path.exists():
            with open(env_path) as f:
                for line in f:
                    if line.startswith("OPENAI_API_KEY="):
                        api_key = line.split("=", 1)[1].strip().strip('"').strip("'")
                        break
    
    if not api_key:
        print("ERROR: OPENAI_API_KEY not found")
        return 1
    
    # Create evaluator
    config = LLMConfig(
        provider="openai",
        model=args.model,
        api_key=api_key,
        temperature=0.3
    )
    evaluator = LLMEvaluator(config)
    
    print("=" * 60)
    print("EXCEL EVALUATION AGENT")
    print(f"Model: {args.model}")
    print(f"File: {args.file}")
    print(f"API delay: {API_DELAY}s")
    print("=" * 60)
    
    # Run evaluation
    segments = [args.segment] if args.segment else None
    run_evaluation(args.file, evaluator, segments, args.start_row, args)
    
    print("\n" + "=" * 60)
    print("EVALUATION COMPLETE")
    print("=" * 60)
    
    return 0


if __name__ == "__main__":
    sys.exit(main())

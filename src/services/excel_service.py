"""
Excel Service
===============
Handles all Excel file operations.
Single Responsibility: Only Excel read/write operations.
"""

from pathlib import Path
from typing import Optional, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

from ..models import Question, Response, Evaluation
from ..config import paths, EXCEL_COLUMNS, EXCEL_HEADERS


class ExcelService:
    """Service for Excel file operations."""
    
    def __init__(self, excel_path: Optional[Path] = None):
        self.excel_path = excel_path or paths.excel_file
        self._workbook: Optional[Workbook] = None
        self._worksheet = None
    
    def open(self) -> bool:
        """Open the Excel file."""
        try:
            self._workbook = load_workbook(self.excel_path)
            self._worksheet = self._workbook.active
            return True
        except Exception as e:
            print(f"Error opening Excel: {e}")
            return False
    
    def save(self) -> bool:
        """Save the Excel file."""
        if self._workbook:
            try:
                self._workbook.save(self.excel_path)
                return True
            except PermissionError:
                print("Error: Excel file is open. Please close it.")
                return False
        return False
    
    def close(self):
        """Close the workbook."""
        if self._workbook:
            self._workbook.close()
            self._workbook = None
            self._worksheet = None
    
    def ensure_headers(self):
        """Ensure all headers are properly set."""
        if not self._worksheet:
            return
        
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        
        for col, header in EXCEL_HEADERS.items():
            cell = self._worksheet[f"{col}1"]
            if cell.value != header:
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
        
        # Set column widths
        self._set_column_widths()
    
    def _set_column_widths(self):
        """Set appropriate column widths."""
        width_map = {
            "A": 50,  # Question
            # Response columns: 50, Eval: 12, Comment: 40
        }
        
        for col in "ABCDEFGHIJKLMNOPQRSTUV":
            if col in ["A", "B", "I", "P"]:  # Question & Response columns
                self._worksheet.column_dimensions[col].width = 50
            elif col in ["C", "E", "G", "J", "L", "N", "Q", "S", "U"]:  # Eval columns
                self._worksheet.column_dimensions[col].width = 12
            else:  # Comment columns
                self._worksheet.column_dimensions[col].width = 40
    
    def get_questions(self, start_row: int = 2) -> List[Tuple[int, str]]:
        """
        Get all questions from column A.
        
        Returns:
            List of (row_number, question_text) tuples
        """
        if not self._worksheet:
            return []
        
        questions = []
        for row in range(start_row, self._worksheet.max_row + 1):
            text = self._worksheet[f"A{row}"].value
            if text and str(text).strip():
                questions.append((row, str(text)))
        
        return questions
    
    def get_response(self, segment: str, row: int) -> Optional[str]:
        """Get response for a segment at a specific row."""
        if segment not in EXCEL_COLUMNS:
            return None
        
        response_col = EXCEL_COLUMNS[segment][0]
        cell_value = self._worksheet[f"{response_col}{row}"].value
        return str(cell_value) if cell_value else None
    
    def get_evaluation(self, segment: str, row: int) -> Optional[Tuple[int, str]]:
        """Get existing evaluation (score, comment) for a segment at a row."""
        if segment not in EXCEL_COLUMNS:
            return None
        
        _, eval_col, comment_col = EXCEL_COLUMNS[segment]
        
        score = self._worksheet[f"{eval_col}{row}"].value
        comment = self._worksheet[f"{comment_col}{row}"].value
        
        if score is not None and score != "":
            try:
                score_int = int(score)
                return (score_int, str(comment) if comment else "")
            except (ValueError, TypeError):
                # Score is not a valid integer
                return None
        return None
    
    def write_evaluation(self, segment: str, row: int, score: int, comment: str = "") -> bool:
        """
        Write evaluation result to Excel.
        
        Args:
            segment: Segment name
            row: Row number
            score: Score (0-5)
            comment: Comment (only written if score <= 3)
        """
        if segment not in EXCEL_COLUMNS:
            return False
        
        _, eval_col, comment_col = EXCEL_COLUMNS[segment]
        
        self._worksheet[f"{eval_col}{row}"] = score
        
        # Only write comment if score <= 3
        if score <= 3:
            self._worksheet[f"{comment_col}{row}"] = comment
        else:
            self._worksheet[f"{comment_col}{row}"] = ""
        
        return True
    
    def write_response(self, segment: str, row: int, response: str) -> bool:
        """Write a response to Excel."""
        if segment not in EXCEL_COLUMNS:
            return False
        
        response_col = EXCEL_COLUMNS[segment][0]
        self._worksheet[f"{response_col}{row}"] = response
        return True
    
    def has_response(self, segment: str, row: int) -> bool:
        """Check if a response exists."""
        response = self.get_response(segment, row)
        return bool(response and response.strip())
    
    def has_evaluation(self, segment: str, row: int) -> bool:
        """Check if an evaluation exists."""
        return self.get_evaluation(segment, row) is not None

